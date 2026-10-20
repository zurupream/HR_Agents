"""Pytest suite for T015: Human in the loop — HR decides uncertain cases and can always override."""

from __future__ import annotations

import asyncio
import json
from pathlib import Path
import sys

from google.adk.runners import InMemoryRunner
from google.genai import types
import openpyxl
import pytest

REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

import screening_agent.agent as agent_module
from screening_agent.agent import apply_uncertainty_rules
from screening_agent.schemas import (
    CandidateEvaluation,
    Criterion,
    CriterionScore,
    JobRubric,
)
from screening_agent.tests.fixtures.create_dummy_data import build_dummy_dataset
import screening_agent.tools as tools_module
from screening_agent.tools import (
    HIGH_REVIEW_RATE_WARNING,
    compute_candidate_status,
    record_hr_decision,
    write_results_excel,
)


@pytest.fixture(autouse=True)
def _sales_test_dataset(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Path:
    """Create temporary Sales Executive fixture dataset in tmp_path for every test."""
    data_dir = tmp_path / "sales_fixture"
    out_dir = tmp_path / "sales_output"
    out_dir.mkdir(parents=True, exist_ok=True)
    build_dummy_dataset(data_dir)
    monkeypatch.setattr(tools_module, "DEFAULT_ADVERT_PATH", data_dir / "job_advert.docx")
    monkeypatch.setattr(tools_module, "DEFAULT_APPLICATIONS_DIR", data_dir / "applications")
    monkeypatch.setattr(tools_module, "DEFAULT_RUBRIC_PATH", out_dir / "rubric.json")
    monkeypatch.setattr(tools_module, "DEFAULT_EXCEL_PATH", out_dir / "screening_results.xlsx")
    monkeypatch.setattr(tools_module, "DEFAULT_EVALUATIONS_PATH", out_dir / "evaluations.json")
    monkeypatch.setattr(tools_module, "DEFAULT_HR_DECISIONS_PATH", out_dir / "hr_decisions.json")
    return data_dir


def _sample_rubric(shortlist_threshold: float = 70.0) -> JobRubric:
    return JobRubric(
        job_title="Sales Executive - Retail and Business Clients",
        role_summary="Drive B2B and retail sales revenue, manage CRM pipelines, and close commercial deals.",
        must_have_criteria=[
            Criterion(
                id="M1",
                name="B2B & Retail Sales Experience with Targets",
                category="must_have",
                description="3+ years B2B or retail sales experience exceeding targets.",
                weight=40.0,
            ),
            Criterion(
                id="M2",
                name="Closing Deals & Pipeline Conversion",
                category="must_have",
                description="Full-cycle deal closing and commercial conversion.",
                weight=35.0,
            ),
        ],
        nice_to_have_criteria=[
            Criterion(
                id="N1",
                name="Negotiation & Client Presentations",
                category="nice_to_have",
                description="Commercial negotiation and client presentations.",
                weight=25.0,
            ),
        ],
        shortlist_threshold=shortlist_threshold,
        requires_continuity=False,
    )


def _sample_mapping() -> dict[str, dict[str, str]]:
    return {
        "C001": {"name": "Dummy Candidate A", "folder": "Dummy Candidate A"},
        "C002": {"name": "Dummy Candidate B", "folder": "Dummy Candidate B"},
        "C003": {"name": "Dummy Candidate C", "folder": "Dummy Candidate C"},
        "C004": {"name": "Dummy Candidate D", "folder": "Dummy Candidate D"},
    }


def test_a_clear_candidate_gets_ai_decided_and_no_human_review() -> None:
    """(a) A clear candidate gets status 'AI decided' and needs_human_review False."""
    rubric = _sample_rubric(shortlist_threshold=70.0)
    clear_eval = CandidateEvaluation(
        candidate_id="C001",
        criterion_scores=[
            CriterionScore(
                criterion_id="M1",
                criterion_name="B2B & Retail Sales Experience with Targets",
                category="must_have",
                weight=40.0,
                score=5,
                evidence='"6 years of B2B and retail sales exceeding revenue targets."',
            ),
            CriterionScore(
                criterion_id="M2",
                criterion_name="Closing Deals & Pipeline Conversion",
                category="must_have",
                weight=35.0,
                score=5,
                evidence='"Closed over 140 commercial and retail accounts."',
            ),
            CriterionScore(
                criterion_id="N1",
                criterion_name="Negotiation & Client Presentations",
                category="nice_to_have",
                weight=25.0,
                score=4,
                evidence='"Delivered executive commercial presentations and contract negotiations."',
            ),
        ],
        overall_score=95.0,
        decision="Shortlisted",
        confidence=0.95,
        review_reasons=[],
        strengths=[
            "B2B & Retail Sales Experience with Targets (5/5)",
            "Closing Deals & Pipeline Conversion (5/5)",
        ],
        gaps=["No significant gaps identified"],
        reasoning="Candidate C001 is a strong match across all sales criteria.",
        needs_human_review=False,
    )

    evaluated = apply_uncertainty_rules(
        clear_eval,
        rubric,
        app_text="=== CV ===\n6 years of B2B and retail sales exceeding revenue targets and closing deals.",
    )
    assert evaluated.needs_human_review is False
    assert evaluated.review_reasons == []
    assert compute_candidate_status(evaluated) == "AI decided"


def test_b_borderline_score_gets_borderline_score_and_needs_hr_review() -> None:
    """(b) A borderline score (within 10 of the threshold) gets 'borderline_score' and status 'Needs HR review'."""
    rubric = _sample_rubric(shortlist_threshold=70.0)
    borderline_eval = CandidateEvaluation(
        candidate_id="C002",
        criterion_scores=[
            CriterionScore(
                criterion_id="M1",
                criterion_name="B2B & Retail Sales Experience with Targets",
                category="must_have",
                weight=40.0,
                score=3,
                evidence='"2 years of retail showroom sales meeting store targets."',
            ),
            CriterionScore(
                criterion_id="M2",
                criterion_name="Closing Deals & Pipeline Conversion",
                category="must_have",
                weight=35.0,
                score=3,
                evidence='"Assisted senior account executives with lead qualification."',
            ),
            CriterionScore(
                criterion_id="N1",
                criterion_name="Negotiation & Client Presentations",
                category="nice_to_have",
                weight=25.0,
                score=4,
                evidence='"Presented product pricing to walk-in clients."',
            ),
        ],
        overall_score=65.0,  # within 10 points of 70.0 threshold
        decision="Not Shortlisted",
        confidence=0.72,
        review_reasons=[],
        strengths=["Negotiation & Client Presentations (4/5)"],
        gaps=["B2B & Retail Sales Experience with Targets (3/5)"],
        reasoning="Candidate C002 is close to the threshold at 65.0/100.",
        needs_human_review=False,
    )

    evaluated = apply_uncertainty_rules(
        borderline_eval,
        rubric,
        app_text="=== CV ===\n2 years of retail showroom sales and client presentations.",
    )
    assert "borderline_score" in evaluated.review_reasons
    assert evaluated.needs_human_review is True
    assert compute_candidate_status(evaluated) == "Needs HR review"


def test_c_missing_evidence_for_must_have_gets_missing_evidence() -> None:
    """(c) Missing evidence for a must-have gets 'missing_evidence'."""
    rubric = _sample_rubric(shortlist_threshold=70.0)
    missing_must_have_eval = CandidateEvaluation(
        candidate_id="C004",
        criterion_scores=[
            CriterionScore(
                criterion_id="M1",
                criterion_name="B2B & Retail Sales Experience with Targets",
                category="must_have",
                weight=40.0,
                score=5,
                evidence='"5 years of B2B and retail sales exceeding revenue targets."',
            ),
            CriterionScore(
                criterion_id="M2",
                criterion_name="Closing Deals & Pipeline Conversion",
                category="must_have",
                weight=35.0,
                score=0,
                evidence="No supporting evidence found in the application documents.",
            ),
            CriterionScore(
                criterion_id="N1",
                criterion_name="Negotiation & Client Presentations",
                category="nice_to_have",
                weight=25.0,
                score=3,
                evidence='"Conducted client presentations for regional accounts."',
            ),
        ],
        overall_score=55.0,  # >10 away from 70.0, so only missing_evidence triggers
        decision="Not Shortlisted",
        confidence=0.75,
        review_reasons=[],
        strengths=["B2B & Retail Sales Experience with Targets (5/5)"],
        gaps=["Closing Deals & Pipeline Conversion (0/5)"],
        reasoning="Strong sales background but no evidence for must-have M2.",
        needs_human_review=False,
    )

    evaluated = apply_uncertainty_rules(
        missing_must_have_eval,
        rubric,
        app_text="=== CV ===\n5 years of B2B and retail sales exceeding revenue targets.",
    )
    assert "missing_evidence" in evaluated.review_reasons
    assert evaluated.needs_human_review is True
    assert compute_candidate_status(evaluated) == "Needs HR review"


def test_d_record_hr_decision_writes_json_regenerates_excel_and_sets_final_decision(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """(d) record_hr_decision writes hr_decisions.json, regenerates the Excel, and Final Decision equals the HR decision."""
    rubric_path = tmp_path / "rubric.json"
    evals_path = tmp_path / "evaluations.json"
    hr_path = tmp_path / "hr_decisions.json"
    excel_path = tmp_path / "screening_results.xlsx"

    monkeypatch.setattr(tools_module, "DEFAULT_RUBRIC_PATH", rubric_path)
    monkeypatch.setattr(tools_module, "DEFAULT_EVALUATIONS_PATH", evals_path)
    monkeypatch.setattr(tools_module, "DEFAULT_HR_DECISIONS_PATH", hr_path)
    monkeypatch.setattr(tools_module, "DEFAULT_EXCEL_PATH", excel_path)
    monkeypatch.setattr(tools_module, "_CANDIDATE_MAPPING", _sample_mapping())

    rubric = _sample_rubric(shortlist_threshold=70.0)
    rubric_path.write_text(rubric.model_dump_json(indent=2), encoding="utf-8")

    evaluations = [
        CandidateEvaluation(
            candidate_id="C001",
            criterion_scores=[],
            overall_score=95.0,
            decision="Shortlisted",
            confidence=0.95,
            review_reasons=[],
            strengths=["B2B & Retail Sales (5/5)"],
            gaps=["None"],
            reasoning="Clear strong match.",
            needs_human_review=False,
        ),
        CandidateEvaluation(
            candidate_id="C004",
            criterion_scores=[],
            overall_score=75.0,
            decision="Not Shortlisted",
            confidence=0.68,
            review_reasons=["missing_evidence"],
            strengths=["B2B & Retail Sales (5/5)"],
            gaps=["CRM Software Proficiency (0/5)"],
            reasoning="Missing evidence on must-have M3.",
            needs_human_review=True,
        ),
    ]
    write_results_excel(
        evaluations=evaluations,
        rubric=rubric,
        candidate_mapping=_sample_mapping(),
        output_path=excel_path,
        hr_decisions={},
    )

    # Record HR decision for C004
    result = record_hr_decision(
        candidate_id="C004",
        decision="Shortlisted",
        note="Verified Salesforce CRM experience during initial phone check",
    )
    assert result["candidate_id"] == "C004"
    assert result["decision"] == "Shortlisted"
    assert result["status"] == "HR decided"
    assert hr_path.exists()

    hr_json = json.loads(hr_path.read_text(encoding="utf-8"))
    assert "C004" in hr_json
    assert hr_json["C004"]["candidate_id"] == "C004"
    assert hr_json["C004"]["decision"] == "Shortlisted"
    assert (
        hr_json["C004"]["note"]
        == "Verified Salesforce CRM experience during initial phone check"
    )
    assert "timestamp" in hr_json["C004"]

    # Check regenerated Excel
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Results"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    assert headers == [
        "Candidate ID",
        "Name",
        "AI Recommendation",
        "AI Confidence",
        "Review Reasons",
        "Status",
        "HR Decision",
        "HR Note",
        "Final Decision",
        "Overall Score",
        "Top Strengths",
        "Gaps",
        "Reasoning",
        "Source Folder",
    ]

    rows_by_cid = {}
    for r in range(2, ws.max_row + 1):
        row_dict = {
            headers[c - 1]: ws.cell(row=r, column=c).value
            for c in range(1, len(headers) + 1)
        }
        rows_by_cid[row_dict["Candidate ID"]] = row_dict

    assert rows_by_cid["C004"]["AI Recommendation"] == "Not Shortlisted"
    assert rows_by_cid["C004"]["Status"] == "HR decided"
    assert rows_by_cid["C004"]["HR Decision"] == "Shortlisted"
    assert (
        rows_by_cid["C004"]["HR Note"]
        == "Verified Salesforce CRM experience during initial phone check"
    )
    assert rows_by_cid["C004"]["Final Decision"] == "Shortlisted"


def test_e_hr_decision_overrides_ai_recommendation_in_both_directions(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """(e) An HR decision overrides an AI recommendation in both directions."""
    rubric_path = tmp_path / "rubric.json"
    evals_path = tmp_path / "evaluations.json"
    hr_path = tmp_path / "hr_decisions.json"
    excel_path = tmp_path / "screening_results.xlsx"

    monkeypatch.setattr(tools_module, "DEFAULT_RUBRIC_PATH", rubric_path)
    monkeypatch.setattr(tools_module, "DEFAULT_EVALUATIONS_PATH", evals_path)
    monkeypatch.setattr(tools_module, "DEFAULT_HR_DECISIONS_PATH", hr_path)
    monkeypatch.setattr(tools_module, "DEFAULT_EXCEL_PATH", excel_path)
    monkeypatch.setattr(tools_module, "_CANDIDATE_MAPPING", _sample_mapping())

    rubric = _sample_rubric(shortlist_threshold=70.0)
    rubric_path.write_text(rubric.model_dump_json(indent=2), encoding="utf-8")

    evaluations = [
        CandidateEvaluation(
            candidate_id="C001",
            criterion_scores=[],
            overall_score=96.0,
            decision="Shortlisted",  # AI recommends Shortlisted
            confidence=0.95,
            review_reasons=[],
            strengths=["All criteria (5/5)"],
            gaps=["None"],
            reasoning="AI shortlisted C001.",
            needs_human_review=False,
        ),
        CandidateEvaluation(
            candidate_id="C002",
            criterion_scores=[],
            overall_score=52.0,
            decision="Not Shortlisted",  # AI recommends Not Shortlisted
            confidence=0.84,
            review_reasons=[],
            strengths=["Retail Sales (3/5)"],
            gaps=["Deal Closing (2/5)"],
            reasoning="AI did not shortlist C002.",
            needs_human_review=False,
        ),
    ]
    write_results_excel(
        evaluations=evaluations,
        rubric=rubric,
        candidate_mapping=_sample_mapping(),
        output_path=excel_path,
        hr_decisions={},
    )

    # Direction 1: AI Shortlisted -> HR overrides to Not Shortlisted
    record_hr_decision(
        candidate_id="C001",
        decision="Not Shortlisted",
        note="Withdrew application / salary mismatch",
    )
    # Direction 2: AI Not Shortlisted -> HR overrides to Shortlisted
    record_hr_decision(
        candidate_id="C002",
        decision="Shortlisted",
        note="strong portfolio",
    )

    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Results"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    rows_by_cid = {}
    for r in range(2, ws.max_row + 1):
        row_dict = {
            headers[c - 1]: ws.cell(row=r, column=c).value
            for c in range(1, len(headers) + 1)
        }
        rows_by_cid[row_dict["Candidate ID"]] = row_dict

    assert rows_by_cid["C001"]["AI Recommendation"] == "Shortlisted"
    assert rows_by_cid["C001"]["HR Decision"] == "Not Shortlisted"
    assert rows_by_cid["C001"]["Final Decision"] == "Not Shortlisted"
    assert rows_by_cid["C001"]["Status"] == "HR decided"

    assert rows_by_cid["C002"]["AI Recommendation"] == "Not Shortlisted"
    assert rows_by_cid["C002"]["HR Decision"] == "Shortlisted"
    assert rows_by_cid["C002"]["Final Decision"] == "Shortlisted"
    assert rows_by_cid["C002"]["Status"] == "HR decided"


def test_f_review_rate_warning_appears_when_more_than_40_percent_need_review(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """(f) The review-rate warning appears when more than 40% need review."""
    evals_path = tmp_path / "evaluations.json"
    excel_path = tmp_path / "screening_results.xlsx"
    monkeypatch.setattr(tools_module, "DEFAULT_EVALUATIONS_PATH", evals_path)

    rubric = _sample_rubric(shortlist_threshold=70.0)
    # 2 out of 3 candidates (66.7% > 40%) need HR review
    evaluations = [
        CandidateEvaluation(
            candidate_id="C001",
            criterion_scores=[],
            overall_score=95.0,
            decision="Shortlisted",
            confidence=0.95,
            review_reasons=[],
            strengths=["All criteria (5/5)"],
            gaps=["None"],
            reasoning="Clear match.",
            needs_human_review=False,
        ),
        CandidateEvaluation(
            candidate_id="C002",
            criterion_scores=[],
            overall_score=66.0,
            decision="Not Shortlisted",
            confidence=0.68,
            review_reasons=["borderline_score"],
            strengths=["Retail Sales (3/5)"],
            gaps=["Deal Closing (2/5)"],
            reasoning="Borderline score.",
            needs_human_review=True,
        ),
        CandidateEvaluation(
            candidate_id="C004",
            criterion_scores=[],
            overall_score=74.0,
            decision="Not Shortlisted",
            confidence=0.65,
            review_reasons=["missing_evidence", "borderline_score"],
            strengths=["B2B & Retail Sales (5/5)"],
            gaps=["CRM Software Proficiency (0/5)"],
            reasoning="Missing must-have evidence.",
            needs_human_review=True,
        ),
    ]

    write_results_excel(
        evaluations=evaluations,
        rubric=rubric,
        candidate_mapping=_sample_mapping(),
        output_path=excel_path,
        hr_decisions={},
    )

    wb = openpyxl.load_workbook(excel_path)
    assert "Review Queue" in wb.sheetnames
    assert "Metadata" in wb.sheetnames

    # Verify sorting in Results: Needs HR review first, then by score descending (C004 74.0, C002 66.0, then C001 95.0)
    ws_results = wb["Results"]
    result_cids = [ws_results.cell(row=r, column=1).value for r in range(2, ws_results.max_row + 1)]
    assert result_cids == ["C004", "C002", "C001"]

    # Verify Review Queue sheet contains only Needs HR review candidates (C004, C002)
    ws_queue = wb["Review Queue"]
    queue_cids = [ws_queue.cell(row=r, column=1).value for r in range(2, ws_queue.max_row + 1)]
    assert queue_cids == ["C004", "C002"]

    # Verify Metadata sheet contains the exact calibration warning
    ws_meta = wb["Metadata"]
    meta_dict = {
        ws_meta.cell(row=r, column=1).value: ws_meta.cell(row=r, column=2).value
        for r in range(2, ws_meta.max_row + 1)
    }
    assert int(meta_dict["AI Decided"]) == 1
    assert int(meta_dict["Needs HR Review"]) == 2
    assert int(meta_dict["HR Decided"]) == 0
    assert "66.7%" in str(meta_dict["Flag Rate"])
    assert meta_dict["Warning"] == HIGH_REVIEW_RATE_WARNING

    # Also verify the final message emitted by the pipeline includes the warning when > 40% need review
    monkeypatch.delenv("GOOGLE_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    monkeypatch.delenv("GOOGLE_GENAI_USE_VERTEXAI", raising=False)
    monkeypatch.setattr(tools_module, "DEFAULT_RUBRIC_PATH", tmp_path / "rubric.json")
    monkeypatch.setattr(tools_module, "DEFAULT_HR_DECISIONS_PATH", tmp_path / "hr_decisions.json")
    monkeypatch.setattr(tools_module, "DEFAULT_EXCEL_PATH", excel_path)

    def _fake_evaluate_high_flag(self: object, full_prompt_text: str) -> CandidateEvaluation:
        import re as _re
        m = _re.search(r"CANDIDATE_ID\s*:\s*(C\d{3})", full_prompt_text)
        cid = m.group(1) if m else "C001"
        return CandidateEvaluation(
            candidate_id=cid,
            criterion_scores=[],
            overall_score=68.0,
            decision="Not Shortlisted",
            confidence=0.55,
            review_reasons=["borderline_score", "low_confidence"],
            strengths=["Retail Sales (3/5)"],
            gaps=["Deal Closing (2/5)"],
            reasoning="Uncertain borderline candidate.",
            needs_human_review=True,
        )

    monkeypatch.setattr(
        agent_module.ScreeningAdkLlm, "_evaluate_candidate", _fake_evaluate_high_flag
    )
    pipeline_out = agent_module.run_screening_pipeline()
    assert pipeline_out["summary"].get("warning") == HIGH_REVIEW_RATE_WARNING
    assert HIGH_REVIEW_RATE_WARNING in json.dumps(pipeline_out["summary"])


def test_adk_web_chat_invokes_record_hr_decision(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """Verify HR can say 'Shortlist C002, note: strong portfolio' in ADK chat and record_hr_decision runs."""
    monkeypatch.delenv("GOOGLE_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    monkeypatch.delenv("GOOGLE_GENAI_USE_VERTEXAI", raising=False)

    rubric_path = tmp_path / "rubric.json"
    evals_path = tmp_path / "evaluations.json"
    hr_path = tmp_path / "hr_decisions.json"
    excel_path = tmp_path / "screening_results.xlsx"

    monkeypatch.setattr(tools_module, "DEFAULT_RUBRIC_PATH", rubric_path)
    monkeypatch.setattr(tools_module, "DEFAULT_EVALUATIONS_PATH", evals_path)
    monkeypatch.setattr(tools_module, "DEFAULT_HR_DECISIONS_PATH", hr_path)
    monkeypatch.setattr(tools_module, "DEFAULT_EXCEL_PATH", excel_path)

    async def _run_chat() -> None:
        runner = InMemoryRunner(
            agent=agent_module.root_agent,
            app_name="adk_web_hitl_test",
        )
        session = await runner.session_service.create_session(
            app_name="adk_web_hitl_test",
            user_id="hr_user",
            state={},
        )

        # Turn 1: initial screening via "hi"
        msg1 = types.Content(role="user", parts=[types.Part.from_text(text="hi")])
        async for _ in runner.run_async(
            user_id="hr_user", session_id=session.id, new_message=msg1
        ):
            pass

        # Turn 2: HR overrides C002 via chat
        msg2 = types.Content(
            role="user",
            parts=[types.Part.from_text(text="Shortlist C002, note: strong portfolio")],
        )
        async for _ in runner.run_async(
            user_id="hr_user", session_id=session.id, new_message=msg2
        ):
            pass

        assert hr_path.exists()
        decisions = json.loads(hr_path.read_text(encoding="utf-8"))
        assert decisions["C002"]["decision"] == "Shortlisted"
        assert decisions["C002"]["note"] == "strong portfolio"

    asyncio.run(_run_chat())
