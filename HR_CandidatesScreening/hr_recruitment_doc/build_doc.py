"""Build the AgentWorks HR & Recruitment Use Case Word document (T011)."""

from __future__ import annotations

from pathlib import Path
from docx import Document

BASE_DIR = Path(__file__).resolve().parent
OUTPUT_PATH = BASE_DIR / "output" / "AgentWorks_HR_Recruitment_Use_Case.docx"


def build_use_case_document(output_path: Path = OUTPUT_PATH) -> Path:
    """Create AgentWorks_HR_Recruitment_Use_Case.docx with exact required content and structure."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    # Title
    doc.add_heading("AI Agents for HR & Recruitment", level=0)

    # Intro
    doc.add_paragraph(
        "Hiring managers drown in inbound resumes, scheduling ping-pong, and onboarding checklists "
        "- while candidates expect fast, fair communication."
    )

    # Section: The challenge
    doc.add_heading("The challenge", level=1)
    challenge_bullets = [
        "Screening against structured criteria takes hours per role, and unconscious bias creeps in when process is informal.",
        "Interview coordination across time zones creates email chains that delay offers and frustrate top candidates.",
        "New hires miss critical policy acknowledgements because onboarding tasks scatter across HRIS, email, and shared drives.",
    ]
    for item in challenge_bullets:
        doc.add_paragraph(item, style="List Bullet")

    # Section: The AgentWorks solution
    doc.add_heading("The AgentWorks solution", level=1)
    solution_paragraphs = [
        "AgentWorks helps HR teams encode rubrics and knockout questions so screening agents summarize fit against explicit criteria - never a black-box hire/no-hire. Humans remain the decision makers for shortlists and offers; agents handle repetitive reading and scheduling logistics.",
        "Interview scheduler agents propose slots that respect panel availability, time zones, and room resources, then sync updates when plans change. Outreach templates stay on-brand and localized while still allowing recruiter edits before send.",
        "Onboarding agents walk new hires through tasks with links to the right policy PDFs and e-sign steps, nudging owners when something stalls. Sensitive topics (compensation, medical, disciplinary) can be excluded from automation by policy.",
        "Activity is centralized for audits and DEI reporting: you can show who saw which candidate summary and when approvals happened. That is increasingly important as regulators scrutinize automated employment tools alongside data protection rules.",
    ]
    for para in solution_paragraphs:
        doc.add_paragraph(para)

    # Section: Workflow in four steps
    doc.add_heading("Workflow in four steps", level=1)
    workflow_steps = [
        "1. Define criteria - Upload job scorecards and policy boundaries so screening stays aligned with legal and DEI guidelines.",
        "2. Screen & summarize - Agents cluster applicants, highlight evidence against rubrics, and flag missing information.",
        "3. Schedule & communicate - Coordinators receive proposed slots and candidate comms drafts ready for one-click approval.",
        "4. Onboard with guardrails - Day-one tasks, policy reads, and system access requests flow as tracked checklists with human oversight.",
    ]
    for step in workflow_steps:
        doc.add_paragraph(step, style="List Number")

    # Section: Worked example - Screening 80 applicants against a scorecard
    doc.add_heading("Worked example - Screening 80 applicants against a scorecard", level=1)
    worked_example_bullets = [
        "Agent: Candidate Screening Agent, then Interview Scheduler Agent",
        "Input: A job scorecard with weighted must-haves and knockout questions, plus the resume and application batch for the requisition.",
        "Approval gate: The agent clusters candidates and summarizes evidence against each scorecard line - it never issues a hire/no-hire verdict. A recruiter reviews the shortlist and manually approves who advances before the Interview Scheduler Agent proposes panel slots.",
        "Output: A ranked shortlist with evidence citations per criterion, plus (once approved) proposed interview slots synced to panel calendars and time zones.",
    ]
    for item in worked_example_bullets:
        doc.add_paragraph(item, style="List Bullet")

    # Section: Results & benefits
    doc.add_heading("Results & benefits", level=1)
    results_bullets = [
        "Faster time-to-interview: Less manual reading and calendar tennis means qualified candidates hear back sooner.",
        "Consistent process: Every requisition follows the same documented steps - easier for audits and hiring manager training.",
        "Human-owned decisions: Automation prepares evidence; people retain control on shortlists, messaging, and offers.",
    ]
    for item in results_bullets:
        doc.add_paragraph(item, style="List Bullet")

    doc.save(str(output_path))
    return output_path


if __name__ == "__main__":
    out = build_use_case_document()
    print(f"Created Word document at: {out}")
