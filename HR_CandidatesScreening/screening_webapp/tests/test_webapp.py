"""Pytest suite for the HR Candidate Screening Web Application (T016).

Covers:
1. Auth: no token -> 401; invalid token -> 401; valid token (mocked verifier) for an allowed email -> 200;
   allowed-list empty -> 403; non-allowed email -> 403.
2. Runs: starting a run returns a run_id; a second run while one is running returns 409;
   results are stored in the SQLite database (runs & candidate_results).
3. Decisions: POST decision writes hr_decisions and audit_log, and regenerates the Excel workbook;
   the latest decision wins in GET /api/runs/latest/results while history is preserved.
4. Plain-language reasons: each internal review code (borderline_score, missing_evidence,
   conflicting_evidence, unreadable_document, low_confidence) maps to a readable sentence.
5. Settings validation: model must be in the allow-list; applications_dir outside the project is rejected;
   redact_before_scoring cannot be turned off.
6. End-to-end demo flow: sign in (mocked token), see candidates & CV preview, run screening,
   see plain-language results, record a decision, and verify it in History and exported Excel.
"""

from __future__ import annotations

import json
from pathlib import Path
import threading
import time
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from screening_agent import tools as tools_module
from screening_agent.schemas import ALLOWED_REVIEW_REASONS
from screening_agent.tests.fixtures.create_dummy_data import build_dummy_dataset
from screening_webapp import app as webapp_module
from screening_webapp import auth as auth_module
from screening_webapp import database as db_module
from screening_webapp.database import (
    AllowedDomain,
    AllowedUser,
    AuditLog,
    CandidateResult,
    HrDecision,
    Job,
    Run,
    get_session,
)
from screening_webapp.plain_language import format_review_reason, format_review_reasons


@pytest.fixture()
def test_env(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    """Configure isolated SQLite database, settings file, Sales fixture, and allow-list for each test."""
    db_file = tmp_path / "test_app.db"
    settings_file = tmp_path / "settings.json"
    excel_file = tmp_path / "screening_results.xlsx"
    allowed_file = tmp_path / "allowed_emails.json"
    sales_dataset = build_dummy_dataset(tmp_path / "sales_fixture")

    # Default to empty allow-list (nobody allowed unless configured)
    allowed_file.write_text(json.dumps({"allowed_emails": []}), encoding="utf-8")

    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{db_file}")
    monkeypatch.setenv("ALLOWED_EMAILS_FILE", str(allowed_file))
    monkeypatch.delenv("ALLOWED_EMAILS", raising=False)
    monkeypatch.delenv("DEMO_TEST_AUTH", raising=False)
    monkeypatch.delenv("GOOGLE_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    monkeypatch.delenv("GOOGLE_GENAI_USE_VERTEXAI", raising=False)

    monkeypatch.setattr(webapp_module, "SETTINGS_FILE_PATH", settings_file)
    monkeypatch.setattr(webapp_module, "WEBAPP_EXCEL_PATH", excel_file)
    monkeypatch.setattr(webapp_module, "DEFAULT_APPLICATIONS_DIR", sales_dataset["applications_dir"])
    monkeypatch.setattr(webapp_module, "DEFAULT_ADVERT_PATH", sales_dataset["advert_path"])
    monkeypatch.setattr(tools_module, "DEFAULT_APPLICATIONS_DIR", sales_dataset["applications_dir"])
    monkeypatch.setattr(tools_module, "DEFAULT_ADVERT_PATH", sales_dataset["advert_path"])
    monkeypatch.setattr(tools_module, "DEFAULT_RUBRIC_PATH", tmp_path / "rubric.json")
    monkeypatch.setattr(tools_module, "DEFAULT_EXCEL_PATH", tmp_path / "agent_results.xlsx")
    monkeypatch.setattr(tools_module, "DEFAULT_HR_DECISIONS_PATH", tmp_path / "hr_decisions.json")
    monkeypatch.setattr(tools_module, "DEFAULT_EVALUATIONS_PATH", tmp_path / "evaluations.json")

    # Reset cached engine so it binds to the test DATABASE_URL
    db_module._ENGINE = None
    db_module._ENGINE_URL = None
    db_module._SESSION_LOCAL = None
    db_module.init_db()

    webapp_module._ACTIVE_RUN_ID = None
    webapp_module._RUN_PROGRESS.clear()

    client = TestClient(webapp_module.create_webapp())
    return {
        "client": client,
        "db_file": db_file,
        "settings_file": settings_file,
        "excel_file": excel_file,
        "allowed_file": allowed_file,
        "sales_dataset": sales_dataset,
    }


def _allow_emails(allowed_file: Path, emails: list[str]) -> None:
    allowed_file.write_text(json.dumps({"allowed_emails": emails}), encoding="utf-8")


def _mock_valid_verifier(monkeypatch: pytest.MonkeyPatch, email: str = "hr.manager@company.com") -> None:
    def _fake_verify(raw_token: str) -> dict[str, Any]:
        token = raw_token.replace("Bearer ", "").replace("bearer ", "").strip()
        if token == "valid-google-token":
            return {
                "email": email,
                "name": "HR Manager",
                "sub": "google-sub-123",
                "aud": auth_module.GOOGLE_CLIENT_ID,
                "iss": "https://accounts.google.com",
            }
        raise ValueError("Invalid Google ID token")

    monkeypatch.setattr(auth_module, "verify_google_token", _fake_verify)


# =============================================================================
# 1. Authentication & Allow-list Tests
# =============================================================================


def test_auth_no_token_returns_401(test_env: dict[str, Any]) -> None:
    client: TestClient = test_env["client"]
    # Public routes work without token
    assert client.get("/health").status_code == 200
    assert client.get("/").status_code == 200

    # Protected API routes reject missing token with 401
    assert client.get("/api/folders").status_code == 401
    assert client.get("/api/settings").status_code == 401
    assert client.post("/api/runs").status_code == 401
    assert client.get("/api/runs/latest/results").status_code == 401


def test_auth_invalid_token_returns_401(
    test_env: dict[str, Any], monkeypatch: pytest.MonkeyPatch
) -> None:
    client: TestClient = test_env["client"]
    _allow_emails(test_env["allowed_file"], ["hr.manager@company.com"])
    _mock_valid_verifier(monkeypatch, email="hr.manager@company.com")

    resp = client.get(
        "/api/folders",
        headers={"Authorization": "Bearer wrong-or-expired-token"},
    )
    assert resp.status_code == 401


def test_auth_empty_allowlist_returns_403(
    test_env: dict[str, Any], monkeypatch: pytest.MonkeyPatch
) -> None:
    client: TestClient = test_env["client"]
    # Default allow-list is empty -> nobody can use the app (403)
    _mock_valid_verifier(monkeypatch, email="hr.manager@company.com")

    resp = client.get(
        "/api/folders",
        headers={"Authorization": "Bearer valid-google-token"},
    )
    assert resp.status_code == 403
    assert "empty" in resp.json()["detail"].lower()


def test_auth_valid_token_allowed_email_returns_200(
    test_env: dict[str, Any], monkeypatch: pytest.MonkeyPatch
) -> None:
    client: TestClient = test_env["client"]
    _allow_emails(test_env["allowed_file"], ["hr.manager@company.com"])
    _mock_valid_verifier(monkeypatch, email="hr.manager@company.com")

    resp = client.get(
        "/api/folders",
        headers={"Authorization": "Bearer valid-google-token"},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert len(data["folders"]) == 4

    # Non-allowed email still receives 403
    _mock_valid_verifier(monkeypatch, email="intruder@other.com")
    resp_forbidden = client.get(
        "/api/folders",
        headers={"Authorization": "Bearer valid-google-token"},
    )
    assert resp_forbidden.status_code == 403


def test_demo_allowlist_and_helpful_403_message(
    test_env: dict[str, Any], monkeypatch: pytest.MonkeyPatch
) -> None:
    """T017: Verified token for email in demo_allowed_emails.json -> 200; unlisted email -> 403 with helpful message."""
    client: TestClient = test_env["client"]

    # Production allow-list must remain empty
    prod_path = auth_module.DEFAULT_ALLOWED_EMAILS_PATH
    prod_data = json.loads(prod_path.read_text(encoding="utf-8"))
    assert prod_data.get("allowed_emails") == []

    # Point ALLOWED_EMAILS_FILE to demo_allowed_emails.json using Windows backslash style
    monkeypatch.setenv("ALLOWED_EMAILS_FILE", r"screening_webapp\config\demo_allowed_emails.json")

    # Verified token for guruprem.saginadham@sybyl.com (in demo list) -> 200
    _mock_valid_verifier(monkeypatch, email="guruprem.saginadham@sybyl.com")
    resp_ok = client.get(
        "/api/me",
        headers={"Authorization": "Bearer valid-google-token"},
    )
    assert resp_ok.status_code == 200
    assert resp_ok.json()["email"] == "guruprem.saginadham@sybyl.com"

    # Verified token for email NOT in the demo list -> 403 with helpful message
    _mock_valid_verifier(monkeypatch, email="unlisted.person@sybyl.com")
    resp_denied = client.get(
        "/api/me",
        headers={"Authorization": "Bearer valid-google-token"},
    )
    assert resp_denied.status_code == 403
    detail = resp_denied.json()["detail"]
    assert "unlisted.person@sybyl.com" in detail
    assert "signed in" in detail.lower()
    assert "demo_allowed_emails.json" in detail
    # Must NOT leak other emails in the allow-list or internal directory paths
    assert "guruprem.saginadham@sybyl.com" not in detail
    assert "hr.demo@example.com" not in detail
    assert "screening_webapp/config" not in detail
    assert "screening_webapp\\config" not in detail


# =============================================================================
# 2. Plain-Language Reason Mapping Tests
# =============================================================================


def test_plain_language_reason_mappings() -> None:
    # Exact example from T016 specification
    borderline_sentence = format_review_reason("borderline_score", score=64, threshold=70)
    assert (
        borderline_sentence
        == "Borderline: score is 64 out of 100, 6 points below the threshold of 70."
    )

    # Every allowed internal review code maps to a clear human sentence without raw snake_case codes
    all_sentences = format_review_reasons(
        list(ALLOWED_REVIEW_REASONS),
        score=64.0,
        threshold=70.0,
        confidence=0.52,
        evidence={
            "criterion_scores": [
                {
                    "criterion_id": "M3",
                    "criterion_name": "Pipeline Management & Target Achievement",
                    "category": "must_have",
                    "score": 0,
                }
            ]
        },
    )
    assert len(all_sentences) == len(ALLOWED_REVIEW_REASONS)
    for code, sentence in zip(ALLOWED_REVIEW_REASONS, all_sentences):
        assert isinstance(sentence, str) and len(sentence) > 20
        assert sentence.endswith(".")
        assert code not in sentence  # No raw internal code shown to user


# =============================================================================
# 3. Settings Validation Tests
# =============================================================================


def test_settings_validation(
    test_env: dict[str, Any], monkeypatch: pytest.MonkeyPatch
) -> None:
    client: TestClient = test_env["client"]
    _allow_emails(test_env["allowed_file"], ["hr.manager@company.com"])
    _mock_valid_verifier(monkeypatch)
    headers = {"Authorization": "Bearer valid-google-token"}

    # GET default settings
    get_resp = client.get("/api/settings", headers=headers)
    assert get_resp.status_code == 200
    defaults = get_resp.json()
    assert defaults["threshold"] == 70.0
    assert defaults["requires_continuity"] is False
    assert defaults["model"] == "gemini-4-argon"
    assert defaults["redact_before_scoring"] is True

    # Invalid model not in allow-list -> rejected (400)
    bad_model_resp = client.put(
        "/api/settings",
        headers=headers,
        json={"threshold": 70, "requires_continuity": False, "model": "gpt-4o-unknown"},
    )
    assert bad_model_resp.status_code == 400

    # applications_dir outside project -> rejected (400)
    bad_dir_resp = client.put(
        "/api/settings",
        headers=headers,
        json={
            "threshold": 70,
            "requires_continuity": False,
            "model": "gemini-4-argon",
            "applications_dir": "/etc",
        },
    )
    assert bad_dir_resp.status_code == 400

    bad_traversal_resp = client.put(
        "/api/settings",
        headers=headers,
        json={
            "threshold": 70,
            "requires_continuity": False,
            "model": "gemini-4-argon",
            "applications_dir": "../../outside_project",
        },
    )
    assert bad_traversal_resp.status_code == 400

    # redact_before_scoring cannot be turned off -> rejected (400)
    bad_redact_resp = client.put(
        "/api/settings",
        headers=headers,
        json={
            "threshold": 70,
            "requires_continuity": False,
            "model": "gemini-4-argon",
            "redact_before_scoring": False,
        },
    )
    assert bad_redact_resp.status_code == 400

    # Valid update -> 200
    ok_resp = client.put(
        "/api/settings",
        headers=headers,
        json={
            "threshold": 75,
            "requires_continuity": True,
            "model": "gemini-2.5-pro",
        },
    )
    assert ok_resp.status_code == 200
    updated = ok_resp.json()
    assert updated["threshold"] == 75.0
    assert updated["requires_continuity"] is True
    assert updated["model"] == "gemini-2.5-pro"
    assert updated["redact_before_scoring"] is True


# =============================================================================
# 4. Runs & Concurrency (409) & Database Persistence Tests
# =============================================================================


def test_runs_concurrency_409_and_db_storage(
    test_env: dict[str, Any], monkeypatch: pytest.MonkeyPatch
) -> None:
    client: TestClient = test_env["client"]
    _allow_emails(test_env["allowed_file"], ["hr.manager@company.com"])
    _mock_valid_verifier(monkeypatch)
    headers = {"Authorization": "Bearer valid-google-token"}

    # Hold the background pipeline with an Event so we can deterministically verify 409 on a second run
    gate = threading.Event()
    orig_pipeline = webapp_module.run_screening_pipeline

    def _gated_pipeline() -> dict[str, Any]:
        gate.wait(timeout=5.0)
        return orig_pipeline()

    monkeypatch.setattr(webapp_module, "run_screening_pipeline", _gated_pipeline)

    # Start first run -> returns run_id
    r1 = client.post("/api/runs", headers=headers)
    assert r1.status_code == 200
    run_id = r1.json()["run_id"]
    assert run_id.startswith("run_")

    # Second run while first is still running -> returns 409 Conflict
    r2 = client.post("/api/runs", headers=headers)
    assert r2.status_code == 409

    # Release the gate and wait for background run to finish
    gate.set()
    for _ in range(50):
        st = client.get(f"/api/runs/{run_id}", headers=headers).json()
        if st["status"] in ("finished", "failed"):
            break
        time.sleep(0.1)

    final_status = client.get(f"/api/runs/{run_id}", headers=headers).json()
    assert final_status["status"] == "finished"
    assert final_status["current_candidate_index"] == 4
    assert final_status["total_candidates"] == 4

    # Verify runs and candidate_results are stored in SQLite
    db = get_session()
    try:
        run_row = db.query(Run).filter(Run.id == run_id).first()
        assert run_row is not None
        assert run_row.status == "finished"
        assert run_row.candidate_count == 4

        cand_rows = (
            db.query(CandidateResult).filter(CandidateResult.run_id == run_id).all()
        )
        assert len(cand_rows) == 4
    finally:
        db.close()


# =============================================================================
# 5. Decisions, Audit Log, Excel Regeneration, and End-to-End Demo Flow
# =============================================================================


def test_decisions_override_audit_excel_and_e2e_flow(
    test_env: dict[str, Any], monkeypatch: pytest.MonkeyPatch
) -> None:
    client: TestClient = test_env["client"]
    _allow_emails(test_env["allowed_file"], ["hr.demo@example.com"])
    monkeypatch.setenv("DEMO_TEST_AUTH", "1")
    headers = {"Authorization": "Bearer demo-hr-token"}

    # 1. Sign in & check /api/me
    me_resp = client.get("/api/me", headers=headers)
    assert me_resp.status_code == 200
    assert me_resp.json()["email"] == "hr.demo@example.com"
    assert me_resp.json()["demo_mode"] is True

    # 2. See candidate folders and CV preview
    folders_resp = client.get("/api/folders", headers=headers)
    assert folders_resp.status_code == 200
    folders = folders_resp.json()["folders"]
    assert len(folders) == 4
    assert folders[0]["candidate_id"] == "C001"
    assert any(f["parse_status"] == "OK" for f in folders[0]["files"])

    preview_resp = client.get("/api/folders/C001/preview", headers=headers)
    assert preview_resp.status_code == 200
    preview_data = preview_resp.json()
    assert preview_data["candidate_id"] == "C001"
    assert len(preview_data["files"]) >= 1
    assert all(len(f["preview"]) <= 1500 for f in preview_data["files"])

    # 3. Run screening and wait for completion
    start_resp = client.post("/api/runs", headers=headers)
    assert start_resp.status_code == 200
    run_id = start_resp.json()["run_id"]

    for _ in range(50):
        st = client.get(f"/api/runs/{run_id}", headers=headers).json()
        if st["status"] in ("finished", "failed"):
            break
        time.sleep(0.1)

    # 4. Get plain-language results sorted with "Needs your review" first
    res_resp = client.get("/api/runs/latest/results", headers=headers)
    assert res_resp.status_code == 200
    res_data = res_resp.json()
    results = res_data["results"]
    assert len(results) == 4

    # First row must be the "Needs your review" candidate (C004)
    first_cand = results[0]
    assert first_cand["candidate_id"] == "C004"
    assert first_cand["status"] == "Needs HR review"
    assert first_cand["status_badge"] == "Needs your review"
    assert any("Borderline:" in s for s in first_cand["review_reasons_plain"])
    assert any("Missing evidence:" in s for s in first_cand["review_reasons_plain"])

    # 5. Record an initial HR decision for C004, then a second decision to verify latest wins
    d1_resp = client.post(
        "/api/candidates/C004/decision",
        headers=headers,
        json={"decision": "Not Shortlisted", "note": "Initial hold pending phone call"},
    )
    assert d1_resp.status_code == 200

    d2_resp = client.post(
        "/api/candidates/C004/decision",
        headers=headers,
        json={
            "decision": "Shortlisted",
            "note": "Confirmed B2B sales quota achievement and Salesforce CRM experience in screening call",
        },
    )
    assert d2_resp.status_code == 200
    d2_data = d2_resp.json()
    assert d2_data["decision"] == "Shortlisted"
    assert d2_data["status"] == "HR decided"
    assert d2_data["status_badge"] == "You decided"
    assert Path(d2_data["excel_path"]).exists()

    # 6. Verify latest decision wins in GET /api/runs/latest/results
    updated_results = client.get("/api/runs/latest/results", headers=headers).json()["results"]
    c004_row = next(r for r in updated_results if r["candidate_id"] == "C004")
    assert c004_row["status"] == "HR decided"
    assert c004_row["status_badge"] == "You decided"
    assert c004_row["final_decision"] == "Shortlisted"
    assert c004_row["hr_decision"]["decision"] == "Shortlisted"
    assert (
        c004_row["hr_decision"]["note"]
        == "Confirmed B2B sales quota achievement and Salesforce CRM experience in screening call"
    )
    assert c004_row["hr_decision"]["decided_by_email"] == "hr.demo@example.com"

    # 7. Verify SQLite history (both decisions kept in hr_decisions, audit_log populated)
    hist_resp = client.get("/api/history", headers=headers)
    assert hist_resp.status_code == 200
    hist_data = hist_resp.json()
    assert len(hist_data["runs"]) == 1
    c004_decisions = [
        d for d in hist_data["hr_decisions"] if d["candidate_id"] == "C004"
    ]
    assert len(c004_decisions) == 2
    assert c004_decisions[0]["decision"] == "Shortlisted"
    assert c004_decisions[1]["decision"] == "Not Shortlisted"
    assert any(a["action"] == "hr_decision" for a in hist_data["audit_log"])

    # 8. Verify exported Excel file reflects the latest HR decision
    xlsx_resp = client.get("/api/export.xlsx", headers=headers)
    assert xlsx_resp.status_code == 200
    wb = load_workbook(test_env["excel_file"])
    ws = wb["Results"]
    headers_row = [cell.value for cell in ws[1]]
    cid_col = headers_row.index("Candidate ID") + 1
    status_col = headers_row.index("Status") + 1
    hr_dec_col = headers_row.index("HR Decision") + 1
    final_col = headers_row.index("Final Decision") + 1

    found_c004 = False
    for row_idx in range(2, ws.max_row + 1):
        if ws.cell(row=row_idx, column=cid_col).value == "C004":
            found_c004 = True
            assert ws.cell(row=row_idx, column=status_col).value == "HR decided"
            assert ws.cell(row=row_idx, column=hr_dec_col).value == "Shortlisted"
            assert ws.cell(row=row_idx, column=final_col).value == "Shortlisted"
    assert found_c004 is True


# =============================================================================
# 6. T018 Admin Portal, Database Allow-List, Domain Rules & Migration Tests
# =============================================================================


def test_t018_migration_imports_once_and_does_not_duplicate(
    test_env: dict[str, Any], monkeypatch: pytest.MonkeyPatch
) -> None:
    """Migration imports JSON emails once on first start and does not duplicate on second start."""
    monkeypatch.delenv("ALLOWED_EMAILS_FILE", raising=False)
    monkeypatch.delenv("ALLOWED_EMAILS", raising=False)

    # First start / migration
    migrated_first = auth_module.migrate_access_control()
    assert migrated_first["migrated"] is True

    with get_session() as session:
        users = {u.email: u for u in session.query(AllowedUser).all()}
        assert "guruprem.saginadham@sybyl.com" in users
        assert users["guruprem.saginadham@sybyl.com"].role == "admin"
        assert users["guruprem.saginadham@sybyl.com"].active is True
        assert "hr.demo@example.com" not in users

        # Open decision check: allowed_domains must be empty by default
        domains = session.query(AllowedDomain).all()
        assert domains == []

        migration_logs = (
            session.query(AuditLog).filter(AuditLog.action == "migrate_allowlist").all()
        )
        assert len(migration_logs) == 1
        user_count_first = len(users)

    # Second start / migration must be a no-op and not duplicate rows
    migrated_second = auth_module.migrate_access_control()
    assert migrated_second["migrated"] is False

    with get_session() as session:
        users_after = session.query(AllowedUser).all()
        assert len(users_after) == user_count_first
        migration_logs_after = (
            session.query(AuditLog).filter(AuditLog.action == "migrate_allowlist").all()
        )
        assert len(migration_logs_after) == 1


def test_t018_root_admin_signin_and_access_tab(
    test_env: dict[str, Any], monkeypatch: pytest.MonkeyPatch
) -> None:
    """Root admin signs in -> allowed (200), is_admin=True; any user from the OAuth Client signs in -> allowed (200), role='hr'."""
    monkeypatch.delenv("ALLOWED_EMAILS_FILE", raising=False)
    monkeypatch.delenv("ALLOWED_EMAILS", raising=False)
    auth_module.migrate_access_control()

    client: TestClient = test_env["client"]
    _mock_valid_verifier(monkeypatch, email="guruprem.saginadham@sybyl.com")
    headers = {"Authorization": "Bearer valid-google-token"}

    me_resp = client.get("/api/me", headers=headers)
    assert me_resp.status_code == 200
    me_data = me_resp.json()
    assert me_data["email"] == "guruprem.saginadham@sybyl.com"
    assert me_data["role"] == "admin"
    assert me_data["is_admin"] is True
    assert me_data["is_root_admin"] is True

    # Verify index.html contains the Access tab and has no demo banner or demo sign-in button
    html_resp = client.get("/")
    assert html_resp.status_code == 200
    assert 'id="tab-btn-access"' in html_resp.text
    assert 'id="tab-access"' in html_resp.text
    assert "5. Access" in html_resp.text
    assert 'id="demo-banner"' not in html_resp.text
    assert 'id="btn-demo-signin"' not in html_resp.text

    # Root admin can fetch /api/admin/access
    access_resp = client.get("/api/admin/access", headers=headers)
    assert access_resp.status_code == 200
    access_data = access_resp.json()
    assert any(
        u["email"] == "guruprem.saginadham@sybyl.com" and u["is_root_admin"] is True
        for u in access_data["users"]
    )
    assert access_data["domains"] == []

    # Any user authenticated by the Google OAuth Client can sign in as HR by default
    for email in ("tester.employee@sybyl.com", "any.tester@external-org.com"):
        _mock_valid_verifier(monkeypatch, email=email)
        oauth_me = client.get("/api/me", headers=headers)
        assert oauth_me.status_code == 200
        oauth_data = oauth_me.json()
        assert oauth_data["email"] == email
        assert oauth_data["role"] == "hr"
        assert oauth_data["is_admin"] is False
        assert oauth_data["is_root_admin"] is False


def test_t018_active_vs_deactivated_user_and_403_message(
    test_env: dict[str, Any], monkeypatch: pytest.MonkeyPatch
) -> None:
    """Active allowed user -> allowed; deactivated user -> 403 with exact message and no leaked info."""
    monkeypatch.delenv("ALLOWED_EMAILS_FILE", raising=False)
    monkeypatch.delenv("ALLOWED_EMAILS", raising=False)
    auth_module.migrate_access_control()

    client: TestClient = test_env["client"]
    headers = {"Authorization": "Bearer valid-google-token"}

    # 1. Root admin adds a new HR user
    _mock_valid_verifier(monkeypatch, email="guruprem.saginadham@sybyl.com")
    add_resp = client.post(
        "/api/admin/users",
        headers=headers,
        json={"email": "recruiter.jane@company.com", "role": "hr"},
    )
    assert add_resp.status_code == 200
    assert add_resp.json()["email"] == "recruiter.jane@company.com"
    assert add_resp.json()["role"] == "hr"
    assert add_resp.json()["active"] is True

    # 2. Active allowed user signs in -> 200 immediately (no restart)
    _mock_valid_verifier(monkeypatch, email="recruiter.jane@company.com")
    jane_ok = client.get("/api/me", headers=headers)
    assert jane_ok.status_code == 200
    assert jane_ok.json()["email"] == "recruiter.jane@company.com"
    assert jane_ok.json()["role"] == "hr"
    assert jane_ok.json()["is_admin"] is False

    # 3. Root admin deactivates Jane
    _mock_valid_verifier(monkeypatch, email="guruprem.saginadham@sybyl.com")
    deact_resp = client.patch(
        "/api/admin/users/recruiter.jane@company.com",
        headers=headers,
        json={"active": False},
    )
    assert deact_resp.status_code == 200
    assert deact_resp.json()["active"] is False

    # 4. Deactivated user signs in -> 403 with exact required message and no leaked users/domains
    _mock_valid_verifier(monkeypatch, email="recruiter.jane@company.com")
    jane_denied = client.get("/api/me", headers=headers)
    assert jane_denied.status_code == 403
    detail = jane_denied.json()["detail"]
    assert (
        detail
        == "Your Google account signed in, but it is not authorised to use this app. Ask an admin to add your email."
    )
    assert "guruprem.saginadham@sybyl.com" not in detail
    assert "hr.demo@example.com" not in detail


def test_t018_domain_rules_and_non_admin_403(
    test_env: dict[str, Any], monkeypatch: pytest.MonkeyPatch
) -> None:
    """Domain rule: user with listed domain -> allowed; unlisted domain (when ENFORCE_ALLOWLIST=1) -> 403; non-admin calling admin endpoint -> 403."""
    monkeypatch.delenv("ALLOWED_EMAILS_FILE", raising=False)
    monkeypatch.delenv("ALLOWED_EMAILS", raising=False)
    monkeypatch.setenv("ENFORCE_ALLOWLIST", "1")
    auth_module.migrate_access_control()

    client: TestClient = test_env["client"]
    headers = {"Authorization": "Bearer valid-google-token"}

    # 1. User from unlisted domain example.gov -> 403
    _mock_valid_verifier(monkeypatch, email="officer@example.gov")
    unlisted_resp = client.get("/api/me", headers=headers)
    assert unlisted_resp.status_code == 403
    assert (
        unlisted_resp.json()["detail"]
        == "Your Google account signed in, but it is not authorised to use this app. Ask an admin to add your email."
    )

    # 2. Root admin adds domain example.gov (role hr)
    _mock_valid_verifier(monkeypatch, email="guruprem.saginadham@sybyl.com")
    dom_add = client.post(
        "/api/admin/domains",
        headers=headers,
        json={"domain": "example.gov", "role": "hr"},
    )
    assert dom_add.status_code == 200
    assert dom_add.json()["domain"] == "example.gov"
    assert dom_add.json()["role"] == "hr"
    assert dom_add.json()["active"] is True

    # 3. User from listed domain example.gov -> allowed (200) on next request
    _mock_valid_verifier(monkeypatch, email="officer@example.gov")
    listed_resp = client.get("/api/me", headers=headers)
    assert listed_resp.status_code == 200
    assert listed_resp.json()["email"] == "officer@example.gov"
    assert listed_resp.json()["role"] == "hr"
    assert listed_resp.json()["is_admin"] is False

    # 4. Non-admin HR user calls admin endpoints -> 403
    assert client.get("/api/admin/access", headers=headers).status_code == 403
    assert (
        client.post(
            "/api/admin/users",
            headers=headers,
            json={"email": "hacker@example.com", "role": "admin"},
        ).status_code
        == 403
    )
    assert (
        client.post(
            "/api/admin/domains",
            headers=headers,
            json={"domain": "other.org", "role": "hr"},
        ).status_code
        == 403
    )

    # 5. Root admin deactivates domain example.gov -> officer@example.gov gets 403
    _mock_valid_verifier(monkeypatch, email="guruprem.saginadham@sybyl.com")
    dom_deact = client.patch(
        "/api/admin/domains/example.gov",
        headers=headers,
        json={"active": False},
    )
    assert dom_deact.status_code == 200
    assert dom_deact.json()["active"] is False

    _mock_valid_verifier(monkeypatch, email="officer@example.gov")
    assert client.get("/api/me", headers=headers).status_code == 403


def test_t018_last_admin_guard_root_admin_guard_and_audit_log(
    test_env: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """Cannot remove the last active admin -> 409. Root admin cannot be deactivated -> 409. Every change writes an audit_log row."""
    monkeypatch.delenv("ALLOWED_EMAILS_FILE", raising=False)
    monkeypatch.delenv("ALLOWED_EMAILS", raising=False)
    auth_module.migrate_access_control()

    client: TestClient = test_env["client"]
    headers = {"Authorization": "Bearer valid-google-token"}
    _mock_valid_verifier(monkeypatch, email="guruprem.saginadham@sybyl.com")

    # 1. Root admin cannot be deactivated, demoted, or deleted -> 409
    deact_root = client.patch(
        "/api/admin/users/guruprem.saginadham@sybyl.com",
        headers=headers,
        json={"active": False},
    )
    assert deact_root.status_code == 409

    demote_root = client.patch(
        "/api/admin/users/guruprem.saginadham@sybyl.com",
        headers=headers,
        json={"role": "hr"},
    )
    assert demote_root.status_code == 409

    del_root = client.delete(
        "/api/admin/users/guruprem.saginadham@sybyl.com",
        headers=headers,
    )
    assert del_root.status_code == 409

    # 2. Add a second admin (second.admin@company.com), then test last-active-admin guard
    add_admin = client.post(
        "/api/admin/users",
        headers=headers,
        json={"email": "second.admin@company.com", "role": "admin"},
    )
    assert add_admin.status_code == 200

    # Simulate a scenario where second.admin@company.com is the ONLY active admin
    # by pointing ROOT_ADMINS_FILE to an empty list and deactivating guruprem in DB
    empty_roots = tmp_path / "empty_root_admins.json"
    empty_roots.write_text("[]", encoding="utf-8")
    monkeypatch.setenv("ROOT_ADMINS_FILE", str(empty_roots))

    # Sign in as second.admin@company.com and deactivate guruprem (since 2 admins exist, this succeeds)
    _mock_valid_verifier(monkeypatch, email="second.admin@company.com")
    deact_first = client.patch(
        "/api/admin/users/guruprem.saginadham@sybyl.com",
        headers=headers,
        json={"active": False},
    )
    assert deact_first.status_code == 200

    # Now second.admin@company.com is the LAST active admin -> cannot deactivate, demote, or delete (409)
    assert (
        client.patch(
            "/api/admin/users/second.admin@company.com",
            headers=headers,
            json={"active": False},
        ).status_code
        == 409
    )
    assert (
        client.patch(
            "/api/admin/users/second.admin@company.com",
            headers=headers,
            json={"role": "hr"},
        ).status_code
        == 409
    )
    assert (
        client.delete(
            "/api/admin/users/second.admin@company.com",
            headers=headers,
        ).status_code
        == 409
    )

    # 3. Verify every change wrote an audit_log row
    access_data = client.get("/api/admin/access", headers=headers).json()
    audit_actions = [row["action"] for row in access_data["audit_log"]]
    assert "migrate_allowlist" in audit_actions
    assert "access_grant_user" in audit_actions
    assert "access_revoke_user" in audit_actions


# =============================================================================
# 7. T019 Multi-Job Management, Per-Job Runs, Decisions, Excel Export & UX Tests
# =============================================================================


def test_t019_create_jobs_run_filter_decisions_and_export(
    test_env: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """T019: Create two jobs, run both, filter results by job_id, record per-job decisions, and export per-job Excel."""
    client: TestClient = test_env["client"]
    _allow_emails(test_env["allowed_file"], ["guruprem.saginadham@sybyl.com"])
    _mock_valid_verifier(monkeypatch, email="guruprem.saginadham@sybyl.com")
    headers = {"Authorization": "Bearer valid-google-token"}

    sales_dataset = test_env["sales_dataset"]
    jd_bytes = Path(sales_dataset["advert_path"]).read_bytes()
    app_dir = Path(sales_dataset["applications_dir"])

    # 1. Create Job 1 via multipart upload (2 candidate folders: Dummy Candidate A and Dummy Candidate B)
    cand_a_cv = (app_dir / "Dummy Candidate A" / "cv.pdf").read_bytes()
    cand_a_cover = (app_dir / "Dummy Candidate A" / "cover_letter.docx").read_bytes()
    cand_b_cv = (app_dir / "Dummy Candidate B" / "cv.docx").read_bytes()

    j1_resp = client.post(
        "/api/jobs",
        headers=headers,
        data={
            "title": "Sales Executive - Retail Clients",
            "threshold": "70",
            "requires_continuity": "false",
            "model": "gemini-4-argon",
            "cv_paths": [
                "applications/Dummy Candidate A/cv.pdf",
                "applications/Dummy Candidate A/cover_letter.docx",
                "applications/Dummy Candidate B/cv.docx",
            ],
        },
        files=[
            ("jd_file", ("retail_sales_jd.docx", jd_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
            ("cv_files", ("cv.pdf", cand_a_cv, "application/pdf")),
            ("cv_files", ("cover_letter.docx", cand_a_cover, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
            ("cv_files", ("cv.docx", cand_b_cv, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
        ],
    )
    assert j1_resp.status_code == 200
    j1_data = j1_resp.json()
    assert j1_data["job_id"] == "JOB-2026-001"
    assert j1_data["job"]["job_id"] == "JOB-2026-001"
    assert j1_data["title"] == "Sales Executive - Retail Clients"
    assert j1_data["candidate_count"] == 2
    assert Path(j1_data["jd_path"]).exists()
    assert Path(j1_data["applications_dir"]).exists()

    # Verify document count badges on Candidates endpoint for JOB-2026-001
    j1_folders = client.get("/api/jobs/JOB-2026-001/folders", headers=headers).json()["folders"]
    assert len(j1_folders) == 2
    assert j1_folders[0]["document_count"] == 2
    assert j1_folders[0]["documents_label"] == "2 documents"
    assert j1_folders[1]["document_count"] == 1
    assert j1_folders[1]["documents_label"] == "1 document"

    # 2. Create Job 2 pointing to the 4-candidate Sales Executive fixture
    j2_resp = client.post(
        "/api/jobs",
        headers=headers,
        json={
            "title": "Sales Executive - Retail and Business Clients",
            "jd_path": str(sales_dataset["advert_path"]),
            "applications_dir": str(sales_dataset["applications_dir"]),
            "threshold": 70.0,
            "requires_continuity": False,
            "model": "gemini-4-argon",
        },
    )
    assert j2_resp.status_code == 200
    j2_data = j2_resp.json()
    assert j2_data["job_id"] == "JOB-2026-002"
    assert j2_data["candidate_count"] == 4

    # Verify GET /api/jobs lists both jobs and supports search by Job ID
    all_jobs = client.get("/api/jobs", headers=headers).json()["jobs"]
    assert [j["job_id"] for j in all_jobs] == ["JOB-2026-002", "JOB-2026-001"]
    searched_jobs = client.get("/api/jobs?search=JOB-2026-001", headers=headers).json()["jobs"]
    assert len(searched_jobs) == 1
    assert searched_jobs[0]["job_id"] == "JOB-2026-001"

    # 3. Run screening for JOB-2026-001 (2 candidates) and wait for completion
    run1_start = client.post("/api/jobs/JOB-2026-001/runs", headers=headers)
    assert run1_start.status_code == 200
    run1_id = run1_start.json()["run_id"]
    for _ in range(50):
        st = client.get(f"/api/runs/{run1_id}", headers=headers).json()
        if st["status"] in ("finished", "failed"):
            break
        time.sleep(0.1)
    assert client.get(f"/api/runs/{run1_id}", headers=headers).json()["status"] == "finished"

    # 4. Run screening for JOB-2026-002 (4 candidates) and wait for completion
    run2_start = client.post("/api/jobs/JOB-2026-002/runs", headers=headers)
    assert run2_start.status_code == 200
    run2_id = run2_start.json()["run_id"]
    for _ in range(50):
        st = client.get(f"/api/runs/{run2_id}", headers=headers).json()
        if st["status"] in ("finished", "failed"):
            break
        time.sleep(0.1)
    assert client.get(f"/api/runs/{run2_id}", headers=headers).json()["status"] == "finished"

    # 5. Filter Results by Job ID: JOB-2026-001 has 2 candidates, JOB-2026-002 has 4 candidates
    res1 = client.get("/api/jobs/JOB-2026-001/results", headers=headers).json()
    assert res1["job_id"] == "JOB-2026-001"
    assert res1["summary"]["total_candidates"] == 2
    assert len(res1["results"]) == 2
    assert all(r["job_id"] == "JOB-2026-001" for r in res1["results"])

    res2 = client.get("/api/jobs/JOB-2026-002/results", headers=headers).json()
    assert res2["job_id"] == "JOB-2026-002"
    assert res2["summary"]["total_candidates"] == 4
    assert len(res2["results"]) == 4
    assert all(r["job_id"] == "JOB-2026-002" for r in res2["results"])

    # Verify truncated Why preview vs full Why text in results
    assert any(
        r["why_truncated"] is True
        and r["why_short"].endswith("...")
        and len(r["why_full"]) > len(r["why_short"])
        for r in res2["results"]
    )

    # Verify filter ("Needs your review only"), search, sort, and pagination on JOB-2026-002
    res2_review_only = client.get(
        "/api/jobs/JOB-2026-002/results?needs_review_only=true", headers=headers
    ).json()
    assert 0 < len(res2_review_only["results"]) < 4
    assert all(r["status"] == "Needs HR review" for r in res2_review_only["results"])
    assert res2_review_only["results"][0]["candidate_id"] == "C004"

    res2_page1 = client.get(
        "/api/jobs/JOB-2026-002/results?page=1&page_size=2&sort_by=score_desc",
        headers=headers,
    ).json()
    assert res2_page1["pagination"]["total_items"] == 4
    assert res2_page1["pagination"]["total_pages"] == 2
    assert len(res2_page1["results"]) == 2
    assert res2_page1["results"][0]["ai_score"] >= res2_page1["results"][1]["ai_score"]

    # 6. Record HR decisions keyed by (job_id, candidate_id) and verify isolation across jobs
    dec_j1 = client.post(
        "/api/jobs/JOB-2026-001/candidates/C001/decision",
        headers=headers,
        json={"decision": "Not Shortlisted", "note": "Role filled for Retail Clients"},
    )
    assert dec_j1.status_code == 200
    assert dec_j1.json()["job_id"] == "JOB-2026-001"

    dec_j2 = client.post(
        "/api/jobs/JOB-2026-002/candidates/C001/decision",
        headers=headers,
        json={"decision": "Shortlisted", "note": "Strong B2B and retail sales track record"},
    )
    assert dec_j2.status_code == 200
    assert dec_j2.json()["job_id"] == "JOB-2026-002"

    res1_after = client.get("/api/jobs/JOB-2026-001/results", headers=headers).json()["results"]
    c1_j1 = next(r for r in res1_after if r["candidate_id"] == "C001")
    assert c1_j1["final_decision"] == "Not Shortlisted"
    assert c1_j1["hr_decision"]["note"] == "Role filled for Retail Clients"

    res2_after = client.get("/api/jobs/JOB-2026-002/results", headers=headers).json()["results"]
    c1_j2 = next(r for r in res2_after if r["candidate_id"] == "C001")
    assert c1_j2["final_decision"] == "Shortlisted"
    assert c1_j2["hr_decision"]["note"] == "Strong B2B and retail sales track record"

    # 7. Verify per-job Excel exports (screening_results_JOB-2026-001.xlsx & screening_results_JOB-2026-002.xlsx)
    xlsx1 = client.get("/api/jobs/JOB-2026-001/export.xlsx", headers=headers)
    assert xlsx1.status_code == 200
    assert "screening_results_JOB-2026-001.xlsx" in xlsx1.headers.get("content-disposition", "")
    xlsx1_file = tmp_path / "jobs" / "JOB-2026-001" / "screening_results_JOB-2026-001.xlsx"
    assert xlsx1_file.exists()
    wb1 = load_workbook(xlsx1_file)
    assert wb1["Results"].max_row == 3  # header + 2 candidates
    meta1_values = [
        (wb1["Metadata"].cell(row=r, column=1).value, wb1["Metadata"].cell(row=r, column=2).value)
        for r in range(2, wb1["Metadata"].max_row + 1)
    ]
    assert ("Job ID", "JOB-2026-001") in meta1_values

    xlsx2 = client.get("/api/jobs/JOB-2026-002/export.xlsx", headers=headers)
    assert xlsx2.status_code == 200
    assert "screening_results_JOB-2026-002.xlsx" in xlsx2.headers.get("content-disposition", "")
    xlsx2_file = tmp_path / "jobs" / "JOB-2026-002" / "screening_results_JOB-2026-002.xlsx"
    assert xlsx2_file.exists()
    wb2 = load_workbook(xlsx2_file)
    assert wb2["Results"].max_row == 5  # header + 4 candidates
    meta2_values = [
        (wb2["Metadata"].cell(row=r, column=1).value, wb2["Metadata"].cell(row=r, column=2).value)
        for r in range(2, wb2["Metadata"].max_row + 1)
    ]
    assert ("Job ID", "JOB-2026-002") in meta2_values


def test_settings_panel_folded_by_default_and_new_jobs_use_saved_defaults(
    test_env: dict[str, Any], monkeypatch: pytest.MonkeyPatch
) -> None:
    """Verify the Default Screening Settings panel is folded by default, opened from the top bar,
    shows the single-line summary when folded, removes 'Run screening' and 'Active Job / Folder'
    from the settings panel, and still saves defaults used by newly created jobs."""
    client: TestClient = test_env["client"]
    _allow_emails(test_env["allowed_file"], ["guruprem.saginadham@sybyl.com"])
    _mock_valid_verifier(monkeypatch, email="guruprem.saginadham@sybyl.com")
    headers = {"Authorization": "Bearer valid-google-token"}

    html_resp = client.get("/")
    assert html_resp.status_code == 200
    html = html_resp.text

    # 1. Top bar has "Default settings" button next to the theme switch and single-line summary when folded
    assert 'id="btn-default-settings"' in html
    assert 'aria-expanded="false"' in html
    assert 'aria-controls="settings-panel"' in html
    assert "Default settings" in html
    assert 'id="settings-summary"' in html
    assert "Threshold 70 · gemini-4-argon · Redaction on" in html

    # 2. Settings panel is folded away by default
    assert (
        '<aside id="settings-panel" class="settings-panel collapsed hidden" data-folded="true" aria-hidden="true"'
        in html
    )

    # 3. "Run screening" button and "Active Job / Folder" box are removed from settings panel,
    #    and screening runs per job from the Jobs page only
    assert 'id="btn-run-screening"' not in html
    assert 'id="setting-app-dir"' not in html
    assert "Active Job / Folder" not in html
    assert 'id="btn-run-selected-job"' not in html

    # 4. Saving defaults works and new jobs inherit the saved defaults
    save_resp = client.put(
        "/api/settings",
        headers=headers,
        json={
            "threshold": 82,
            "requires_continuity": True,
            "model": "gemini-2.5-pro",
            "redact_before_scoring": True,
        },
    )
    assert save_resp.status_code == 200
    saved = save_resp.json()
    assert saved["threshold"] == 82.0
    assert saved["requires_continuity"] is True
    assert saved["model"] == "gemini-2.5-pro"

    sales_dataset = test_env["sales_dataset"]
    job_resp = client.post(
        "/api/jobs",
        headers=headers,
        json={
            "title": "Enterprise Account Executive",
            "jd_path": str(sales_dataset["advert_path"]),
            "applications_dir": str(sales_dataset["applications_dir"]),
        },
    )
    assert job_resp.status_code == 200
    created_job = job_resp.json()["job"]
    assert created_job["threshold"] == 82.0
    assert created_job["requires_continuity"] is True
    assert created_job["model"] == "gemini-2.5-pro"
    assert created_job["applications_dir"] == str(sales_dataset["applications_dir"])

